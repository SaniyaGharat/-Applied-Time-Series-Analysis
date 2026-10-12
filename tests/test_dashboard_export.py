"""
Unit and Integration Tests for Phase 8 Executive Dashboard Excel Export.
Runnable with:
    pytest tests/test_dashboard_export.py -v
"""

import io
import openpyxl
import pandas as pd
import pytest

from modules.export import build_dashboard_report


def test_build_dashboard_report_openpyxl_loads_and_six_sheets_exist():
    """Verify build_dashboard_report returns valid bytes, openpyxl loads it, and 6 sheets exist."""
    kpi_df = pd.DataFrame({
        "KPI": ["Observations", "Best Model", "Best MASE"],
        "A": ["144", "SARIMA(1,1,1)(1,1,1)12", "0.42"],
        "B": ["100", "ARIMA(1,1,1)", "0.88"],
        "note": ["N", "Winner", "Scale-free"],
    })
    comp_a = pd.DataFrame({
        "Rank": [1, 2],
        "Model": ["SARIMA(1,1,1)(1,1,1)12", "ARIMA(1,1,1)"],
        "MASE": [0.42, 0.55],
    })
    comp_b = pd.DataFrame({
        "Rank": [1, 2],
        "Model": ["ARIMA(1,1,1)", "Naive"],
        "MASE": [0.88, 1.00],
    })
    future_a = pd.DataFrame({
        "mean": [120.0, 125.0, 130.0],
        "lower": [110.0, 112.0, 115.0],
        "upper": [130.0, 138.0, 145.0],
    }, index=pd.date_range("2025-01-01", periods=3, freq="MS"))
    future_b = pd.DataFrame({
        "mean": [50.0, 52.0, 54.0],
        "lower": [45.0, 46.0, 47.0],
        "upper": [55.0, 58.0, 61.0],
    }, index=pd.date_range("2025-01-01", periods=3, freq="MS"))
    findings = [
        "Dataset A displays stronger seasonality than Dataset B.",
        "Model SARIMA won on Dataset A, whereas ARIMA won on Dataset B.",
    ]

    # Generate workbook bytes
    excel_bytes = build_dashboard_report(
        kpi_df=kpi_df,
        comparison_a=comp_a,
        comparison_b=comp_b,
        future_a=future_a,
        future_b=future_b,
        findings=findings,
    )

    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 0

    # Load workbook using openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet_names = wb.sheetnames

    expected_sheets = [
        "KPIs",
        "Findings",
        "Leaderboard_A",
        "Leaderboard_B",
        "Future_A",
        "Future_B",
    ]
    for expected in expected_sheets:
        assert expected in sheet_names, f"Expected sheet '{expected}' missing from workbook sheets {sheet_names}"

    assert len(sheet_names) == 6


def test_build_dashboard_report_none_fallback_sheets():
    """Verify build_dashboard_report gracefully generates all 6 sheets even with empty/None parameters."""
    excel_bytes = build_dashboard_report()
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    expected_sheets = ["KPIs", "Findings", "Leaderboard_A", "Leaderboard_B", "Future_A", "Future_B"]
    for expected in expected_sheets:
        assert expected in wb.sheetnames
