"""
Unit and integration tests for ATSA Phase 5 Model Comparison, Metrics, Forecasting, and Reporting.

Covers:
- interval_coverage known values & winkler_score hand-computed example
- diebold_mariano: identical errors, normal variances test, n<5 guard
- build_comparison_table: column verification, rank order, relative RMSE, failed model handling
- ic_group: ETS, ARIMA differencing groups, naive baseline
- equal_weight_combination: arithmetic forecast mean, absence of intervals
- make_future_exog: repeat_last, repeat_last_season, mean_last_k, linear_trend, and validation errors
- fit_full_and_forecast_detailed & wrapper: dimensions, bounds, DatetimeIndex, SARIMAX validation, non-raising failure
- export: build_excel_report returns bytes with 4 valid sheets loaded via openpyxl
- A3 MASE consistency: mase_m scale independent verification
- A4 test_actual populated in ModelResult
"""

import io
import os
import sys
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import numpy as np
import openpyxl
import pandas as pd
import pytest

from modules.compare import (
    best_model,
    build_comparison_table,
    equal_weight_combination,
    ic_group,
    make_future_exog,
    pairwise_dm,
    validate_uploaded_future_exog,
)
from modules.export import build_excel_report
from modules.metrics import (
    accuracy_table,
    diebold_mariano,
    interval_coverage,
    mae,
    mean_interval_width,
    relative_metric,
    winkler_score,
)
from modules.models import (
    ForecastResult,
    ModelResult,
    fit_full_and_forecast,
    fit_full_and_forecast_detailed,
    fit_model,
    make_model_key,
)


def test_coverage_and_winkler():
    """Test interval_coverage on known values and hand-computed winkler_score."""
    print("\n--- Test interval_coverage & winkler_score ---")
    y_true = [1.0, 2.0, 3.0]
    lo = [0.0, 0.0, 0.0]
    up = [4.0, 4.0, 4.0]
    cov100 = interval_coverage(y_true, lo, up)
    assert cov100 == 100.0, f"Expected 100.0, got {cov100}"

    # One outside [0, 2.5] -> 1 and 2 inside, 3 outside -> 2/3 = 66.666...%
    up_partial = [2.5, 2.5, 2.5]
    cov66 = interval_coverage(y_true, lo, up_partial)
    assert np.isclose(cov66, 200.0 / 3.0), f"Expected ~66.67, got {cov66}"

    # Winkler score hand-computed example:
    # y = [2.0, 0.5, 3.5], lo=[1.0, 1.0, 1.0], up=[3.0, 3.0, 3.0], alpha=0.05
    # y=2.0 inside -> 2.0
    # y=0.5 -> 2.0 + (2/0.05)*(1.0-0.5) = 2.0 + 40*0.5 = 22.0
    # y=3.5 -> 2.0 + (2/0.05)*(3.5-3.0) = 2.0 + 40*0.5 = 22.0
    # mean = (2.0 + 22.0 + 22.0) / 3 = 46.0 / 3 = 15.333333...
    ws = winkler_score([2.0, 0.5, 3.5], [1.0, 1.0, 1.0], [3.0, 3.0, 3.0], alpha=0.05)
    expected_ws = 46.0 / 3.0
    assert np.isclose(ws, expected_ws), f"Expected {expected_ws}, got {ws}"
    print("[PASS] interval_coverage and winkler_score matched exact expectations.")


def test_diebold_mariano():
    """Test Diebold-Mariano test under various error scenarios."""
    print("\n--- Test diebold_mariano ---")
    # Identical error series -> stat nan, note present
    e_same = np.array([1.0, -1.0, 2.0, 0.5, -0.5, 1.2])
    dm_same = diebold_mariano(e_same, e_same, power=2)
    assert np.isnan(dm_same["stat"]), "Identical series should return NaN DM stat"
    assert "Indicative only" in dm_same["note"]

    # e1 ~ N(0,1), e2 ~ N(0,3), n=60, seed 42 -> stat negative (e1 better than e2) and p < 0.05
    rng = np.random.default_rng(42)
    e1 = rng.normal(0, 1.0, 60)
    e2 = rng.normal(0, 3.0, 60)
    dm_diff = diebold_mariano(e1, e2, h=1, power=2)
    assert dm_diff["stat"] < 0, f"Expected negative DM stat (e1 better than e2), got {dm_diff['stat']}"
    assert dm_diff["pvalue"] < 0.05, f"Expected pvalue < 0.05, got {dm_diff['pvalue']}"

    # n < 5 -> returns nan
    dm_short = diebold_mariano([1.0, 2.0, 3.0], [0.5, 1.5, 2.5])
    assert np.isnan(dm_short["stat"])
    assert np.isnan(dm_short["pvalue"])
    print("[PASS] diebold_mariano tests passed.")


def test_pairwise_dm_reference_and_unknown_key():
    """Test pairwise_dm: (a) reference row is the chosen model, (b) unknown key raises ValueError."""
    print("\n--- Test pairwise_dm reference row and unknown key validation ---")
    rng = np.random.default_rng(42)
    dates = pd.date_range("2020-01-01", periods=40, freq="MS")
    s = pd.Series(np.linspace(10, 30, 40) + rng.normal(0, 1, 40), index=dates)

    m1 = fit_model(s, {"family": "ar", "p": 1, "d": 0, "trend": "c"}, test_size=10)
    m2 = fit_model(s, {"family": "arima", "p": 1, "d": 1, "q": 1, "trend": "n"}, test_size=10)

    # (a) Verify reference row is the model chosen by reference_key
    df_m1 = pairwise_dm([m1, m2], reference_key=m1.key)
    ref_rows = df_m1[df_m1["verdict"] == "reference model"]
    assert len(ref_rows) == 1, "Expected exactly 1 reference row"
    assert ref_rows.iloc[0]["Model"] == m1.label, f"Expected reference model '{m1.label}', got '{ref_rows.iloc[0]['Model']}'"

    df_m2 = pairwise_dm([m1, m2], reference_key=m2.key)
    ref_rows2 = df_m2[df_m2["verdict"] == "reference model"]
    assert len(ref_rows2) == 1
    assert ref_rows2.iloc[0]["Model"] == m2.label, f"Expected reference model '{m2.label}', got '{ref_rows2.iloc[0]['Model']}'"

    # (b) Verify unknown key raises ValueError
    with pytest.raises(ValueError, match="not found in results"):
        pairwise_dm([m1, m2], reference_key="nonexistent_key_12345")

    print("[PASS] pairwise_dm reference row and unknown key ValueError verified.")


def test_comparison_table_and_ranking():
    """Test build_comparison_table on simulated series with naive, ar, and arima."""
    print("\n--- Test build_comparison_table & ranking ---")
    rng = np.random.default_rng(42)
    n = 60
    dates = pd.date_range("2020-01-01", periods=n, freq="MS")
    trend = np.linspace(10, 30, n)
    s = pd.Series(trend + rng.normal(0, 1, n), index=dates, name="sales")

    m_naive = fit_model(s, {"family": "naive"}, test_size=12)
    m_ar = fit_model(s, {"family": "ar", "p": 1, "d": 0, "trend": "c"}, test_size=12)
    m_arima = fit_model(s, {"family": "arima", "p": 1, "d": 1, "q": 1, "trend": "n"}, test_size=12)

    # Failed model
    m_failed = fit_model(s, {"family": "arima", "p": 25, "d": 0, "q": 25}, test_size=12)

    table = build_comparison_table([m_naive, m_ar, m_arima, m_failed], rank_by="rmse", reference_family="naive")

    expected_cols = [
        "Model", "Family", "Converged", "RMSE", "MAE", "MAPE", "sMAPE", "MASE",
        "Coverage95", "AvgWidth", "RelRMSE_vs_ref", "AIC", "BIC", "IC_group",
        "LB_min_p", "JB_p", "FitSec", "Rank",
    ]
    for c in expected_cols:
        assert c in table.columns, f"Missing column {c} in comparison table"

    # Naive RelRMSE_vs_ref should be 1.0
    naive_row = table[table["Family"] == "naive"].iloc[0]
    assert np.isclose(naive_row["RelRMSE_vs_ref"], 1.0), f"Expected naive RelRMSE == 1.0, got {naive_row['RelRMSE_vs_ref']}"

    # Rank 1 must have the minimum RMSE among converged models
    conv_table = table[table["Converged"] == True]
    best_row = conv_table[conv_table["Rank"] == 1].iloc[0]
    assert best_row["RMSE"] == conv_table["RMSE"].min()

    # Best model function returns Rank == 1 row's key
    assert best_model(table) == best_row["Key"]
    assert best_model(table, metric="rmse") == best_row["Key"]

    # Ranking by another metric (e.g. MAE)
    table_mae = build_comparison_table([m_naive, m_ar, m_arima, m_failed], rank_by="mae", reference_family="naive")
    best_mae_row = table_mae[table_mae["Rank"] == 1].iloc[0]
    assert best_model(table_mae, metric="MAE") == best_mae_row["Key"]
    assert best_model(table_mae, metric="mae") == best_mae_row["Key"]

    # Failed model is ranked last with NaN rank
    failed_row = table[table["Converged"] == False].iloc[0]
    assert np.isnan(failed_row["Rank"]) or failed_row["Rank"] is None
    print("[PASS] build_comparison_table and best_model assertions passed.")


def test_ic_group():
    """Test ic_group for ARIMA with various d/D, Holt-Winters, and Naive."""
    print("\n--- Test ic_group ---")
    m_hw = ModelResult(
        key="hw", label="HW", family="holt_winters", spec={"family": "holt_winters"},
        train_index=[], test_index=[], fitted_train=pd.Series(), test_forecast=pd.DataFrame(),
        residuals=pd.Series(), metrics_test={}, metrics_train={}, information_criteria={},
        params_table=None, summary_text="", ljung_box=None, normality={},
    )
    assert ic_group(m_hw) == "ETS"

    m_arima_d1 = ModelResult(
        key="ar1", label="AR1", family="arima", spec={"family": "arima", "d": 1, "D": 0, "trend": "n"},
        train_index=[], test_index=[], fitted_train=pd.Series(), test_forecast=pd.DataFrame(),
        residuals=pd.Series(), metrics_test={}, metrics_train={}, information_criteria={},
        params_table=None, summary_text="", ljung_box=None, normality={},
    )
    assert "d=1" in ic_group(m_arima_d1)

    m_arima_d0 = ModelResult(
        key="ar0", label="AR0", family="arima", spec={"family": "arima", "d": 0, "D": 0, "trend": "c"},
        train_index=[], test_index=[], fitted_train=pd.Series(), test_forecast=pd.DataFrame(),
        residuals=pd.Series(), metrics_test={}, metrics_train={}, information_criteria={},
        params_table=None, summary_text="", ljung_box=None, normality={},
    )
    assert ic_group(m_arima_d1) != ic_group(m_arima_d0)

    m_naive = ModelResult(
        key="n", label="Naive", family="naive", spec={"family": "naive"},
        train_index=[], test_index=[], fitted_train=pd.Series(), test_forecast=pd.DataFrame(),
        residuals=pd.Series(), metrics_test={}, metrics_train={}, information_criteria={},
        params_table=None, summary_text="", ljung_box=None, normality={},
    )
    assert ic_group(m_naive) == "n/a"
    print("[PASS] ic_group logic assertions passed.")


def test_equal_weight_combination():
    """Test equal_weight_combination arithmetic mean and interval behavior."""
    print("\n--- Test equal_weight_combination ---")
    idx = pd.date_range("2020-01-01", periods=5, freq="D")
    df1 = pd.DataFrame({"mean": [10.0, 12.0, 14.0, 16.0, 18.0], "lower": [8.0]*5, "upper": [12.0]*5}, index=idx)
    df2 = pd.DataFrame({"mean": [20.0, 22.0, 24.0, 26.0, 28.0], "lower": [18.0]*5, "upper": [22.0]*5}, index=idx)

    actual = pd.Series([15.0, 17.0, 19.0, 21.0, 23.0], index=idx)

    r1 = ModelResult(
        key="m1", label="Model 1", family="ar", spec={"family": "ar"},
        train_index=idx, test_index=idx, fitted_train=pd.Series(index=idx),
        test_forecast=df1, residuals=pd.Series(), metrics_test={}, metrics_train={},
        information_criteria=None, params_table=None, summary_text="", ljung_box=None,
        normality={}, converged=True, test_actual=actual, extras={"mase_scale": 2.0},
    )
    r2 = ModelResult(
        key="m2", label="Model 2", family="arima", spec={"family": "arima"},
        train_index=idx, test_index=idx, fitted_train=pd.Series(index=idx),
        test_forecast=df2, residuals=pd.Series(), metrics_test={}, metrics_train={},
        information_criteria=None, params_table=None, summary_text="", ljung_box=None,
        normality={}, converged=True, test_actual=actual, extras={"mase_scale": 2.0},
    )

    combo = equal_weight_combination([r1, r2])
    assert combo is not None
    expected_mean = np.array([15.0, 17.0, 19.0, 21.0, 23.0])
    assert np.allclose(combo.test_forecast["mean"].values, expected_mean)
    assert np.isnan(combo.test_forecast["lower"]).all()
    assert np.isnan(combo.test_forecast["upper"]).all()
    assert combo.extras.get("mase_scale") == 2.0
    # Mean forecast matches actual exactly, so MAE is 0.0 and MASE is 0.0 / 2.0 = 0.0
    assert np.isclose(combo.metrics_test["mase"], 0.0)
    print("[PASS] equal_weight_combination mean and intervals verified.")


def test_equal_weight_combination_mase_scale_reuse():
    """Test that equal_weight_combination reuses extras['mase_scale'] from members."""
    print("\n--- Test equal_weight_combination MASE scale reuse ---")
    rng = np.random.default_rng(42)
    n = 60
    dates = pd.date_range("2020-01-01", periods=n, freq="MS")
    s = pd.Series(np.linspace(10, 30, n) + rng.normal(0, 1, n), index=dates)

    m1 = fit_model(s, {"family": "ar", "p": 1, "d": 0, "trend": "c"}, test_size=12, mase_m=12)
    m2 = fit_model(s, {"family": "arima", "p": 1, "d": 1, "q": 1, "trend": "n"}, test_size=12, mase_m=12)

    assert "mase_scale" in m1.extras
    assert "mase_scale" in m2.extras
    assert m1.extras["mase_scale"] > 0
    assert np.isclose(m1.extras["mase_scale"], m2.extras["mase_scale"])

    combo = equal_weight_combination([m1, m2])
    assert combo is not None
    assert "mase_scale" in combo.extras
    assert np.isclose(combo.extras["mase_scale"], m1.extras["mase_scale"])
    assert not np.isnan(combo.metrics_test["mase"])
    expected_mase = combo.metrics_test["mae"] / combo.extras["mase_scale"]
    assert np.isclose(combo.metrics_test["mase"], expected_mase)
    print("[PASS] equal_weight_combination successfully reused mase_scale from member models.")


def test_make_future_exog():
    """Test make_future_exog strategies and input validation."""
    print("\n--- Test make_future_exog ---")
    dates = pd.date_range("2020-01-01", periods=8, freq="D")
    exog = pd.DataFrame({"feature1": [1, 2, 3, 4, 1, 2, 3, 4], "feature2": np.arange(8) * 2.0}, index=dates)

    # 1. repeat_last
    f_last = make_future_exog(exog, steps=3, strategy="repeat_last")
    assert len(f_last) == 3
    assert (f_last["feature1"] == 4).all()
    assert (f_last["feature2"] == 14.0).all()

    # 2. repeat_last_season (m=4)
    f_seas = make_future_exog(exog, steps=4, strategy="repeat_last_season", m=4)
    assert list(f_seas["feature1"]) == [1, 2, 3, 4]

    # 3. mean_last_k (k=2)
    f_mean = make_future_exog(exog, steps=2, strategy="mean_last_k", k=2)
    assert np.isclose(f_mean["feature1"].iloc[0], 3.5)

    # 4. linear_trend on feature2 (which has slope 2.0)
    f_trend = make_future_exog(exog, steps=2, strategy="linear_trend")
    assert np.isclose(f_trend["feature2"].iloc[0], 16.0)
    assert np.isclose(f_trend["feature2"].iloc[1], 18.0)

    # Validation errors
    with pytest.raises(ValueError, match="NaN"):
        bad_exog = exog.copy()
        bad_exog.iloc[0, 0] = np.nan
        make_future_exog(bad_exog, steps=2, strategy="repeat_last")

    with pytest.raises(ValueError, match="non-numeric"):
        bad_str_exog = exog.copy()
        bad_str_exog["str_col"] = "invalid"
        make_future_exog(bad_str_exog, steps=2, strategy="repeat_last")

    with pytest.raises(ValueError, match="positive"):
        make_future_exog(exog, steps=0, strategy="repeat_last")
    print("[PASS] make_future_exog strategies and input validation verified.")


def test_fit_full_and_forecast_detailed():
    """Test fit_full_and_forecast_detailed across families and error guards."""
    print("\n--- Test fit_full_and_forecast_detailed ---")
    rng = np.random.default_rng(42)
    dates = pd.date_range("2020-01-01", periods=48, freq="MS")
    s = pd.Series(np.linspace(10, 40, 48) + rng.normal(0, 1, 48), index=dates)

    families = [
        {"family": "ar", "p": 1, "d": 0, "trend": "c"},
        {"family": "arima", "p": 1, "d": 1, "q": 1, "trend": "n"},
        {"family": "sarima", "p": 1, "d": 0, "q": 0, "P": 1, "D": 0, "Q": 0, "m": 12},
        {"family": "holt_winters", "trend": "add", "seasonal": "add", "m": 12},
    ]

    for spec in families:
        res = fit_full_and_forecast_detailed(s, spec, steps=6)
        assert isinstance(res, ForecastResult)
        assert len(res.df) == 6
        assert res.converged is True
        assert (res.df["lower"] <= res.df["mean"]).all()
        assert (res.df["mean"] <= res.df["upper"]).all()
        assert isinstance(res.df.index, pd.DatetimeIndex)

    # SARIMAX without future_exog raises ValueError
    exog_hist = pd.DataFrame({"x": np.arange(48)}, index=dates)
    sarimax_spec = {"family": "sarimax", "p": 1, "d": 0, "q": 0}
    with pytest.raises(ValueError, match="future_exog"):
        fit_full_and_forecast_detailed(s, sarimax_spec, steps=6, exog=exog_hist, future_exog=None)

    # SARIMAX with valid future_exog
    future_exog = pd.DataFrame({"x": np.arange(48, 54)}, index=pd.date_range("2024-01-01", periods=6, freq="MS"))
    res_sx = fit_full_and_forecast_detailed(s, sarimax_spec, steps=6, exog=exog_hist, future_exog=future_exog)
    assert res_sx.converged is True

    # Wrapper returns same df
    df_wrap = fit_full_and_forecast(s, sarimax_spec, steps=6, exog=exog_hist, future_exog=future_exog)
    assert np.allclose(df_wrap["mean"], res_sx.df["mean"])
    print("[PASS] fit_full_and_forecast_detailed tests passed.")


def test_excel_export_openpyxl():
    """Test build_excel_report returns openpyxl-loadable workbook with 4 sheets."""
    print("\n--- Test build_excel_report ---")
    comp_df = pd.DataFrame({"Model": ["AR", "ARIMA"], "RMSE": [1.2, 1.0]})
    holdout_dict = {"AR": pd.DataFrame({"mean": [1.0, 2.0]})}
    future_df = pd.DataFrame({"mean": [3.0, 4.0], "lower": [2.0, 3.0], "upper": [4.0, 5.0]})
    specs = [{"label": "AR", "family": "ar", "spec": {"p": 1}}]

    wb_bytes = build_excel_report(
        comparison_df=comp_df,
        holdout_forecasts=holdout_dict,
        future_forecast_df=future_df,
        specs=specs,
        metadata={"holdout_size": 12, "mase_m": 12},
    )
    assert isinstance(wb_bytes, bytes)
    assert len(wb_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(wb_bytes))
    expected_sheets = ["Comparison", "Holdout_Forecasts", "Future_Forecast", "Specs"]
    for s in expected_sheets:
        assert s in wb.sheetnames, f"Expected sheet '{s}' in workbook, found {wb.sheetnames}"
    print("[PASS] openpyxl multi-sheet workbook generation verified.")


def test_a3_mase_consistency_and_a4_test_actual():
    """A3 & A4: MASE scale independence test and test_actual population."""
    print("\n--- Test A3 MASE consistency & A4 test_actual ---")
    rng = np.random.default_rng(42)
    dates = pd.date_range("2020-01-01", periods=60, freq="MS")
    s = pd.Series(np.linspace(10, 50, 60) + rng.normal(0, 1, 60), index=dates)

    # Fit model with mase_m=12
    m1 = fit_model(s, {"family": "ar", "p": 1, "d": 0, "m": 4}, test_size=12, mase_m=12)
    m2 = fit_model(s, {"family": "ar", "p": 2, "d": 0, "m": 4}, test_size=12, mase_m=12)

    # Both models should share identical MASE denominator scale
    train_y = s.iloc[:-12]
    # In-sample seasonal naive difference for m=12
    scale_12 = np.mean(np.abs(train_y.values[12:] - train_y.values[:-12]))

    assert np.isclose(m1.metrics_test["mase"], m1.metrics_test["mae"] / scale_12)
    assert np.isclose(m2.metrics_test["mase"], m2.metrics_test["mae"] / scale_12)

    # A4 test_actual check
    assert m1.test_actual is not None
    assert len(m1.test_actual) == 12
    assert (m1.test_actual.values == s.iloc[-12:].values).all()
    print("[PASS] A3 MASE scale consistency and A4 test_actual verified.")


if __name__ == "__main__":
    print("==================================================")
    print("Starting ATSA Compare & Forecast Tests...")
    print("==================================================")
    test_coverage_and_winkler()
    test_diebold_mariano()
    test_pairwise_dm_reference_and_unknown_key()
    test_comparison_table_and_ranking()
    test_ic_group()
    test_equal_weight_combination()
    test_equal_weight_combination_mase_scale_reuse()
    test_make_future_exog()
    test_fit_full_and_forecast_detailed()
    test_excel_export_openpyxl()
    test_a3_mase_consistency_and_a4_test_actual()
    print("\n==================================================")
    print("ALL TESTS IN TEST_COMPARE COMPLETED SUCCESSFULLY!")
    print("==================================================")
